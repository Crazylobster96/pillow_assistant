"""Deterministic spreadsheet operations for the excel_inspect / excel_edit tools.

Agents asked to "reorder these rows" or "make the font 宋体 10" used to write a
fresh openpyxl script every time, inspect the sheet from every angle and drift
into unrelated audits. These functions do the common jobs in one call, keep
the formatting (cell styles, row heights, merged cells, validations and
conditional formats outside the touched rows), never overwrite the source
unless asked, and report what changed together with checks of the result.

Everything here is openpyxl and the standard library; nothing runs a model.
"""

from __future__ import annotations

import copy
import os
import shutil
import stat
from collections import Counter
from pathlib import Path
from typing import Any, Optional

MAX_GROUPS = 200
SAMPLE_ROWS = 5


class ExcelError(ValueError):
    pass


def _load(path: Path, read_only: bool = False):
    import openpyxl

    if not path.is_file():
        raise ExcelError(f"file not found: {path}")
    if path.suffix.lower() not in (".xlsx", ".xlsm"):
        raise ExcelError("only .xlsx / .xlsm workbooks are supported (save .xls / .csv as .xlsx first)")
    try:
        return openpyxl.load_workbook(path, read_only=read_only, keep_vba=path.suffix.lower() == ".xlsm")
    except Exception as exc:  # noqa: BLE001 - corrupt file, encrypted workbook, ...
        raise ExcelError(f"cannot open {path.name}: {exc}") from None


def _sheet(wb, name: Optional[str]):
    if not name:
        return wb.worksheets[0]
    if name not in wb.sheetnames:
        raise ExcelError(f"no sheet {name!r}; sheets: {', '.join(wb.sheetnames)}")
    return wb[name]


def dims(ws) -> tuple[int, int]:
    """(max_row, max_column), computed once per sheet.

    openpyxl recomputes both by scanning every cell on each access; inside a
    row loop that made a 19 000-row sheet take minutes.
    """
    cached = getattr(ws, "_pillow_dims", None)
    if cached is None:
        cached = (ws.max_row, ws.max_column)
        ws._pillow_dims = cached
    return cached


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def guess_header_row(ws, limit: int = 15) -> int:
    """First row (within ``limit``) holding three or more text cells, else 1."""
    max_row, max_col = dims(ws)
    for r in range(1, min(max_row, limit) + 1):
        texts = [c for c in range(1, max_col + 1)
                 if isinstance(ws.cell(r, c).value, str) and ws.cell(r, c).value.strip()]
        merged_title = any(r in range(m.min_row, m.max_row + 1) and m.max_col - m.min_col >= 2
                           for m in ws.merged_cells.ranges)
        if len(texts) >= 3 and not merged_title:
            return r
    return 1


def column_index(ws, header_row: int, column) -> int:
    """A column given as a header name, a letter (``C``) or a 1-based number."""
    from openpyxl.utils import column_index_from_string

    if isinstance(column, int):
        return column
    text = _text(column)
    for c in range(1, dims(ws)[1] + 1):
        if _text(ws.cell(header_row, c).value) == text:
            return c
    if text.isdigit():
        return int(text)
    try:
        return column_index_from_string(text.upper())
    except ValueError:
        headers = [_text(ws.cell(header_row, c).value) for c in range(1, dims(ws)[1] + 1)]
        raise ExcelError(f"no column {column!r}; headers: {headers}") from None


def last_data_row(ws, first: int) -> int:
    last = first - 1
    max_row, max_col = dims(ws)
    for r in range(first, max_row + 1):
        if any(ws.cell(r, c).value not in (None, "") for c in range(1, max_col + 1)):
            last = r
    return last


def blocks(ws, col: int, first: int, last: int) -> list[tuple[str, int, int]]:
    """Contiguous runs of equal values in ``col``: (value, first row, last row)."""
    out: list[tuple[str, int, int]] = []
    for r in range(first, last + 1):
        value = _text(ws.cell(r, col).value)
        if out and out[-1][0] == value:
            out[-1] = (value, out[-1][1], r)
        else:
            out.append((value, r, r))
    return out


def inspect(path: str | Path, sheet: Optional[str] = None, group_by=None,
            within: Optional[dict] = None) -> dict:
    """What a task needs to know before editing: sheets, header, columns, blocks."""
    path = Path(path)
    wb = _load(path)
    ws = _sheet(wb, sheet)
    header = guess_header_row(ws)
    first, last = header + 1, last_data_row(ws, header + 1)
    info: dict[str, Any] = {
        "file": str(path), "read_only_flag": not os.access(path, os.W_OK),
        "sheets": wb.sheetnames, "sheet": ws.title, "max_row": dims(ws)[0], "max_column": dims(ws)[1],
        "title": _text(ws.cell(1, 1).value) if header > 1 else "",
        "header_row": header,
        "columns": [_text(ws.cell(header, c).value) for c in range(1, dims(ws)[1] + 1)],
        "data_rows": [first, last],
        "sample": [[ws.cell(r, c).value for c in range(1, dims(ws)[1] + 1)]
                   for r in range(first, min(last, first + SAMPLE_ROWS - 1) + 1)],
        "merged": [str(m) for m in list(ws.merged_cells.ranges)[:20]],
        "validations": [str(d.sqref) for d in ws.data_validations.dataValidation][:20],
        "conditional_formats": [str(r.sqref) for r in ws.conditional_formatting][:20],
        "fonts": Counter((ws.cell(r, 1).font.name, ws.cell(r, 1).font.sz)
                         for r in range(first, last + 1)).most_common(5),
    }
    if group_by is not None:
        lo, hi = first, last
        if within:
            lo, hi = _within_rows(ws, header, first, last, within)
        col = column_index(ws, header, group_by)
        found = blocks(ws, col, lo, hi)
        info["groups"] = {"column": info["columns"][col - 1] or col, "rows": [lo, hi],
                          "blocks": [{"value": v, "rows": [a, b], "count": b - a + 1}
                                     for v, a, b in found[:MAX_GROUPS]],
                          "truncated": len(found) > MAX_GROUPS}
    return info


def _within_rows(ws, header: int, first: int, last: int, within: dict) -> tuple[int, int]:
    """The contiguous rows where column ``within["column"]`` equals ``within["value"]``."""
    col = column_index(ws, header, within.get("column"))
    want = _text(within.get("value"))
    rows = [r for r in range(first, last + 1) if _text(ws.cell(r, col).value) == want]
    if not rows:
        raise ExcelError(f"no rows where {within.get('column')} = {want!r}")
    if rows[-1] - rows[0] + 1 != len(rows):
        raise ExcelError(f"rows where {within.get('column')} = {want!r} are not contiguous; "
                         "sort or filter them first")
    return rows[0], rows[-1]


def _snapshot_row(ws, r: int) -> tuple[list, list, Optional[float]]:
    cells = [ws.cell(r, c) for c in range(1, dims(ws)[1] + 1)]
    return [c.value for c in cells], [copy.copy(c._style) for c in cells], ws.row_dimensions[r].height


def _write_row(ws, r: int, row: tuple[list, list, Optional[float]]) -> None:
    values, styles, height = row
    for c, (value, style) in enumerate(zip(values, styles), start=1):
        cell = ws.cell(r, c)
        cell.value = value
        cell._style = style
    ws.row_dimensions[r].height = height


def _match(value: str, wanted: str, mode: str) -> bool:
    return value == wanted if mode == "exact" else (wanted in value or value in wanted)


def op_reorder_blocks(ws, header: int, first: int, last: int, op: dict) -> dict:
    """Put the blocks of equal ``column`` values into ``order`` (whole rows, styles kept)."""
    order = [str(v).strip() for v in op.get("order") or [] if str(v).strip()]
    if not order:
        raise ExcelError("reorder_blocks needs order: the values in the new order")
    lo, hi = (first, last) if not op.get("within") else _within_rows(ws, header, first, last, op["within"])
    col = column_index(ws, header, op.get("column"))
    mode = "exact" if op.get("match") == "exact" else "contains"
    found = blocks(ws, col, lo, hi)
    values = [v for v, _, _ in found]
    if len(set(values)) != len(values):
        raise ExcelError("some values appear in more than one separate block; sort the rows by "
                         "that column first or narrow the range with within")
    rank = {}
    for v in values:
        hits = [i for i, w in enumerate(order) if _match(v, w, mode)]
        if len(hits) > 1:
            raise ExcelError(f"{v!r} matches several entries of order; use match: exact")
        rank[v] = hits[0] if hits else len(order)  # unlisted blocks keep their place after the listed
    missing = [w for w in order if not any(_match(v, w, mode) for v in values)]
    if missing:
        raise ExcelError(f"not found in that column: {missing}; values there: {values}")
    rows = [_snapshot_row(ws, r) for r in range(lo, hi + 1)]
    labelled = [(rank[_text(row[0][col - 1])], i, row) for i, row in enumerate(rows)]
    labelled.sort(key=lambda item: (item[0], item[1]))
    for offset, (_, _, row) in enumerate(labelled):
        _write_row(ws, lo + offset, row)
    after = [v for v, _, _ in blocks(ws, col, lo, hi)]
    return {"op": "reorder_blocks", "rows": [lo, hi], "before": values, "after": after}


def op_replace(ws, header: int, first: int, last: int, op: dict) -> dict:
    find = str(op.get("find") or "")
    if not find:
        raise ExcelError("replace needs find")
    repl = str(op.get("replace") or "")
    cols = ([column_index(ws, header, c) for c in op["columns"]] if op.get("columns")
            else list(range(1, dims(ws)[1] + 1)))
    whole = bool(op.get("whole_cell"))
    changed, examples = 0, []
    for r in range(first, last + 1):
        for c in cols:
            cell = ws.cell(r, c)
            if not isinstance(cell.value, str):
                continue
            if whole:
                new = repl if cell.value.strip() == find else cell.value
            else:
                new = cell.value.replace(find, repl)
            if new != cell.value:
                if len(examples) < 5:
                    examples.append(f"{cell.coordinate}: {cell.value} → {new}")
                cell.value = new
                changed += 1
    return {"op": "replace", "find": find, "replace": repl, "cells_changed": changed, "examples": examples}


def op_set_font(ws, header: int, first: int, last: int, op: dict) -> dict:
    name, size = op.get("name"), op.get("size")
    if not name and not size:
        raise ExcelError("set_font needs name and/or size")
    top = 1 if op.get("include_header", True) else first
    count = 0
    max_row, max_col = dims(ws)
    for r in range(top, max_row + 1):
        for c in range(1, max_col + 1):
            cell = ws.cell(r, c)
            if cell.value is None and not cell.has_style:
                continue
            font = copy.copy(cell.font)
            if name:
                font.name = name
            if size:
                font.sz = float(size)
            cell.font = font  # bold, colour, underline stay as they were
            count += 1
    return {"op": "set_font", "name": name, "size": size, "cells": count}


def op_renumber(ws, header: int, first: int, last: int, op: dict) -> dict:
    col = column_index(ws, header, op.get("column") or 1)
    start = int(op.get("start", 1))
    key = column_index(ws, header, op["skip_if_blank"]) if op.get("skip_if_blank") else None
    n = start
    for r in range(first, last + 1):
        if key is not None and ws.cell(r, key).value in (None, ""):
            continue
        ws.cell(r, col).value = n
        n += 1
    return {"op": "renumber", "column": col, "from": start, "to": n - 1}


OPS = {"reorder_blocks": op_reorder_blocks, "replace": op_replace,
       "set_font": op_set_font, "renumber": op_renumber}


def default_output(source: Path) -> Path:
    return source.with_name(f"{source.stem}_修改{source.suffix}")


def _make_writable(path: Path) -> None:
    if path.exists() and not os.access(path, os.W_OK):
        os.chmod(path, stat.S_IWRITE | stat.S_IREAD)


def edit(path: str | Path, operations: list[dict], *, sheet: Optional[str] = None,
         output: Optional[str | Path] = None, overwrite: bool = False) -> dict:
    """Apply ``operations`` in order and save; returns a report with checks.

    The result goes to ``output`` (default ``<name>_修改.xlsx`` beside the source).
    ``overwrite=True`` writes over the source after copying it to ``<name>_备份.xlsx``.
    """
    source = Path(path)
    if not operations:
        raise ExcelError("no operations")
    wb = _load(source)
    ws = _sheet(wb, sheet)
    header = guess_header_row(ws)
    first, last = header + 1, last_data_row(ws, header + 1)
    rows_before = last - first + 1
    multiset_before = Counter(tuple(_text(v) for v in _snapshot_row(ws, r)[0]) for r in range(first, last + 1))
    report: dict[str, Any] = {"source": str(source), "sheet": ws.title, "header_row": header,
                              "data_rows": [first, last], "steps": []}
    for op in operations:
        kind = str(op.get("op") or "")
        if kind not in OPS:
            raise ExcelError(f"unknown op {kind!r}; use one of {', '.join(OPS)}")
        report["steps"].append(OPS[kind](ws, header, first, last, op))
    target = Path(output) if output else (source if overwrite else default_output(source))
    if target.resolve() == source.resolve():
        if not overwrite:
            raise ExcelError("output is the source file; pass overwrite: true to replace it")
        backup = source.with_name(f"{source.stem}_备份{source.suffix}")
        if not backup.exists():
            shutil.copy2(source, backup)
        report["backup"] = str(backup)
        _make_writable(source)
    elif target.exists() and not overwrite:
        raise ExcelError(f"{target.name} already exists; choose another output or pass overwrite: true")
    else:
        _make_writable(target)
    try:
        wb.save(target)
    except PermissionError:
        raise ExcelError(f"cannot write {target}: it is open in Excel or another program; close it and retry") from None
    report["output"] = str(target)
    # Check the saved file, not the in-memory workbook.
    check = _sheet(_load(target), ws.title)
    rows_after = last_data_row(check, first) - first + 1
    report["checks"] = {"rows_before": rows_before, "rows_after": rows_after,
                        "row_count_kept": rows_before == rows_after}
    if all(s["op"] in ("reorder_blocks", "set_font") for s in report["steps"]):
        after = Counter(tuple(_text(v) for v in _snapshot_row(check, r)[0]) for r in range(first, first + rows_after))
        report["checks"]["same_rows_only_moved"] = after == multiset_before
    return report
