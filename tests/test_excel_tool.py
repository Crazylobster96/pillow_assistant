"""excel_inspect / excel_edit on a sheet shaped like the 协保员 table:
a merged title, the header on row 3, blocks of counties under each branch."""

from __future__ import annotations

import asyncio
import json
import stat
from pathlib import Path

import openpyxl
import pytest
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation

from pillow_assistant.capabilities import excel_ops
from pillow_assistant.core.tools.base import ToolContext
from pillow_assistant.core.tools.builtin.excel_tool import ExcelEditTool, ExcelInspectTool

COUNTIES = [("合肥中支", "瑶海支公司", 3), ("合肥中支", "肥西支公司", 2), ("合肥中支", "庐江支公司", 2),
            ("合肥中支", "巢湖支公司", 3), ("淮南中支", "淮南中支营业部", 2), ("淮南中支", "寿县支公司", 2)]


def _workbook(path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "协保员统计总表"
    ws.merge_cells("A1:G2")
    rows = [["序号", "中支机构", "县支机构", "服务乡镇", "服务行政村", "协保员姓名", "手机号码"]]
    n = 0
    for branch, county, count in COUNTIES:
        for i in range(count):
            n += 1
            rows.append([n, branch, county, f"{county[:2]}镇", None, f"{county[:2]}{i}", f"139{n:08d}"])
    for r, values in enumerate(rows, start=3):  # header on row 3, under the merged title
        for c, value in enumerate(values, start=1):
            cell = ws.cell(r, c, value)
            if r > 3:
                cell.font = Font(name="宋体", size=11 if values[2] == "巢湖支公司" else 10, bold=(c == 6))
    dv = DataValidation(type="textLength", operator="equal", formula1="11", prompt="请输入11位手机号码")
    ws.add_data_validation(dv)
    dv.add("G20")
    wb.save(path)
    return path


def _column(ws, col: int) -> list:
    return [ws.cell(r, col).value for r in range(4, ws.max_row + 1)]


def test_inspect_finds_header_and_blocks(tmp_path):
    info = excel_ops.inspect(_workbook(tmp_path / "表.xlsx"), group_by="县支机构",
                             within={"column": "中支机构", "value": "合肥中支"})
    assert info["header_row"] == 3 and info["columns"][2] == "县支机构"
    assert info["title"] == "协保员统计总表"
    assert [b["value"] for b in info["groups"]["blocks"]] == ["瑶海支公司", "肥西支公司", "庐江支公司", "巢湖支公司"]
    assert info["groups"]["rows"] == [4, 13]


def test_edit_does_the_four_requests_and_leaves_the_source_alone(tmp_path):
    src = _workbook(tmp_path / "表.xlsx")
    before = src.read_bytes()
    report = excel_ops.edit(src, [
        {"op": "reorder_blocks", "column": "县支机构", "within": {"column": "中支机构", "value": "合肥中支"},
         "order": ["瑶海", "肥西", "巢湖", "庐江"]},
        {"op": "replace", "find": "淮南中支营业部", "replace": "中支营业部", "columns": ["县支机构"]},
        {"op": "set_font", "name": "宋体", "size": 10},
        {"op": "renumber", "column": "序号"},
    ])
    assert src.read_bytes() == before
    out = Path(report["output"])
    assert out.name == "表_修改.xlsx" and report["checks"]["row_count_kept"]
    ws = openpyxl.load_workbook(out).active
    counties = _column(ws, 3)
    assert counties[:10] == ["瑶海支公司"] * 3 + ["肥西支公司"] * 2 + ["巢湖支公司"] * 3 + ["庐江支公司"] * 2
    assert counties[10:12] == ["中支营业部"] * 2 and report["steps"][1]["cells_changed"] == 2
    assert _column(ws, 1) == list(range(1, 15))
    assert {(ws.cell(r, 1).font.name, ws.cell(r, 1).font.sz) for r in range(4, 18)} == {("宋体", 10.0)}
    assert ws.cell(4, 6).font.b  # bold names stay bold
    assert ws.merged_cells.ranges and str(list(ws.merged_cells.ranges)[0]) == "A1:G2"
    assert [str(d.sqref) for d in ws.data_validations.dataValidation] == ["G20"]


def test_pure_reorder_is_checked_as_a_move(tmp_path):
    src = _workbook(tmp_path / "表.xlsx")
    report = excel_ops.edit(src, [{"op": "reorder_blocks", "column": "县支机构", "within": {
        "column": "中支机构", "value": "合肥中支"}, "order": ["巢湖"]}])
    assert report["checks"]["same_rows_only_moved"]
    assert report["steps"][0]["after"][0] == "巢湖支公司"  # listed first, the rest keep their order
    assert report["steps"][0]["after"][1:] == ["瑶海支公司", "肥西支公司", "庐江支公司"]


def test_overwrite_backs_up_and_clears_read_only(tmp_path):
    src = _workbook(tmp_path / "表.xlsx")
    src.chmod(stat.S_IREAD)  # as copied from OneDrive
    report = excel_ops.edit(src, [{"op": "set_font", "name": "黑体"}], overwrite=True)
    assert Path(report["backup"]).name == "表_备份.xlsx" and Path(report["backup"]).exists()
    assert openpyxl.load_workbook(src).active.cell(4, 1).font.name == "黑体"


def test_bad_requests_are_refused_with_a_reason(tmp_path):
    src = _workbook(tmp_path / "表.xlsx")
    with pytest.raises(excel_ops.ExcelError, match="not found"):
        excel_ops.edit(src, [{"op": "reorder_blocks", "column": "县支机构", "within": {
            "column": "中支机构", "value": "合肥中支"}, "order": ["瑶海", "长丰"]}])
    with pytest.raises(excel_ops.ExcelError, match="no column"):
        excel_ops.inspect(src, group_by="不存在的列")
    with pytest.raises(excel_ops.ExcelError, match="already exists"):
        excel_ops.edit(src, [{"op": "set_font", "size": 9}])
        excel_ops.edit(src, [{"op": "set_font", "size": 9}])


def test_tools_take_attachments_and_refuse_other_paths(tmp_path):
    (tmp_path / "desk").mkdir()
    (tmp_path / "ws").mkdir()
    src = _workbook(tmp_path / "desk" / "表.xlsx")
    ctx = ToolContext(workspace=str(tmp_path / "ws"), references=[str(src)])
    result = asyncio.run(ExcelInspectTool()({"path": str(src), "group_by": "中支机构"}, ctx))
    assert result.ok and [b["value"] for b in json.loads(result.text)["groups"]["blocks"]] == ["合肥中支", "淮南中支"]
    result = asyncio.run(ExcelEditTool()({"path": str(src), "operations": [
        {"op": "replace", "find": "淮南中支营业部", "replace": "中支营业部"}], "output": "表_改名.xlsx"}, ctx))
    assert result.ok and (tmp_path / "desk" / "表_改名.xlsx").exists()
    stranger = _workbook(tmp_path / "other.xlsx")
    assert not asyncio.run(ExcelInspectTool()({"path": str(stranger)}, ctx)).ok
